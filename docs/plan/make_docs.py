#!/usr/bin/env python3
"""Build the two Word deliverables for the Tech-Yeti project plan.

- Safal-Sharma-Individual-Project-Plan.docx  (same structure as the leader's plan,
  with the individual Gantt chart)
- Team-Project-Plan-Gantt.docx               (the Gantt chart the leader was missing)
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

FONT = "Times New Roman"

def set_font(run, size=11, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element.rPr.rFonts
    r.set(qn("w:eastAsia"), FONT)

def add_para(doc, text, size=11, bold=False, align=None, space_after=6, justify=True):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if justify:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_font(run, size, bold)
    return p

# ================================================================ individual plan
doc = Document()
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Individual Project Plans and Responsibilities"), 16, True)
p.paragraph_format.space_after = Pt(10)

add_para(doc,
    "Safal Sharma is the developer of Tech-Yeti and holds primary responsibility for the "
    "frontend build, backend API, database schema, and all core system implementation "
    "(proposed share of planned effort: 20%). The team includes Bipin Pija, who serves as "
    "the team leader and testing lead with accountability for sprint planning, system "
    "architecture, client communication, and deployment oversight; Jiban Thapa Magar, who "
    "serves as the deputy and documentation lead; and Sushil Gaihre, who handles 360-degree "
    "UI/UX design and frontend UI.")

add_para(doc, "Safal Sharma \u2013 Development Lead", size=12, bold=True, justify=False, space_after=4)

add_para(doc,
    "As the developer, I am responsible for the frontend build, backend API, database "
    "schema, and all core system implementation. I turn the sprint backlogs defined by the "
    "leader into working software \u2014 authentication, RBAC and account management in Sprint 1; "
    "the 360\u00b0 training engine, quizzes, spot-the-hazard gameplay and scoring in Sprint 2; and "
    "the leaderboard, badges, Safety Passport, retraining trigger and admin analytics in "
    "Sprint 3. I integrate the UI designs from Sushil, deploy every sprint to Replit, and "
    "fix bugs until the project is accomplished.")

rows = [
    ("Weeks 1-2",
     "Attend project kickoff and confirm the developer role; review the client brief and "
     "requirements/SRS draft; set up the development toolchain (Node.js, VS Code, Replit); "
     "research 360\u00b0 web technologies (WebGL/Three.js) and the backend API approach.",
     "Development toolchain confirmed; technology research notes; requirements questions "
     "raised with the leader."),
    ("Weeks 3-5",
     "Set up the development environment with the leader; design the database schema "
     "(users, roles, modules, questions, attempts, scores, badges); create the shared "
     "repository structure and API skeleton; review the use-case diagrams and "
     "specifications for implementation feasibility.",
     "Database schema document; repository with folder structure and API endpoint list; "
     "development environment setup confirmed."),
    ("Weeks 6-7",
     "Implement the authentication API (login, RBAC policy, account management) and the "
     "dashboard API; integrate the login and dashboard UI designed by Sushil; deploy "
     "Sprint 1 to Replit (login, RBAC, account management, dashboards); fix bugs before "
     "the Sprint 1 client presentation; support the presentation and obtain formal sign-off.",
     "Working authentication, RBAC, account management and dashboards; Sprint 1 live "
     "deployment (Replit); client sign-off (W7); client feedback log."),
    ("Weeks 8-9",
     "Implement 360\u00b0 scene integration and the module engine; quiz and spot-the-hazard "
     "gameplay logic; scoring, competency and attempt-history APIs; integrate the 360\u00b0 "
     "scene and activity UI; deploy Sprint 2 to Replit (360\u00b0 scenes, quizzes, "
     "spot-the-hazard, scoring, competency); fix bugs; support the presentation and "
     "obtain formal sign-off.",
     "Working 360\u00b0 scenes, quizzes, spot-the-hazard and scoring; Sprint 2 live "
     "deployment (Replit); client sign-off (W9); client feedback log."),
    ("Weeks 10-11",
     "Implement leaderboard logic, badge awards and the Safety Passport; retraining "
     "trigger (score < 70%); audit log; admin analytics; integrate the leaderboard, "
     "badges and Safety Passport UI; deploy Sprint 3 to Replit; review the backend "
     "implementation with the leader; support the documentation compilation from Jiban; "
     "obtain final sprint sign-off.",
     "Working leaderboard, badges, Safety Passport and analytics; Sprint 3 live "
     "deployment (Replit); client sign-off (W11); documentation reviewed and approved."),
    ("Week 12",
     "Regression-test the complete system; make the final consolidated Replit deployment "
     "(complete, tested, signed-off system); take part in the final system walkthrough; "
     "join the final presentation with all team members; confirm the project handover to "
     "the client and that all assessment submissions are complete.",
     "Final consolidated live deployment; final presentation delivered; project handover "
     "complete; all assessment submissions confirmed."),
]

table = doc.add_table(rows=1, cols=3)
table.style = "Table Grid"
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
for i, htext in enumerate(["Period", "Individual activities", "Expected evidence or output"]):
    hdr[i].text = ""
    run = hdr[i].paragraphs[0].add_run(htext)
    set_font(run, 11, True)
for period, acts, evid in rows:
    cells = table.add_row().cells
    for i, text in enumerate((period, acts, evid)):
        cells[i].text = ""
        run = cells[i].paragraphs[0].add_run(text)
        set_font(run, 11, i == 0)
for row in table.rows:
    row.cells[0].width = Inches(0.9)
    row.cells[1].width = Inches(3.7)
    row.cells[2].width = Inches(2.0)

doc.add_paragraph()
add_para(doc, "Gantt chart", size=12, bold=True, justify=False, space_after=6)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture("gantt-safal.png", width=Inches(6.5))

doc.save("Safal-Sharma-Individual-Project-Plan.docx")
print("wrote Safal-Sharma-Individual-Project-Plan.docx")

# ================================================================ team gantt doc
doc2 = Document()
for section in doc2.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

p = doc2.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
set_font(p.add_run("Tech-Yeti \u2013 Team Project Plan: Gantt Chart (Weeks 1-12)"), 16, True)
p.paragraph_format.space_after = Pt(10)

add_para(doc2,
    "Gantt chart of the 12-week Tech-Yeti project plan, prepared from the team leader's "
    "plan. The work is organised into three sprints \u2014 Sprint 1: Authentication & "
    "Dashboard (Weeks 6-7), Sprint 2: 360\u00b0 Training Core (Weeks 8-9) and Sprint 3: "
    "Engagement & Safety Passport (Weeks 10-11) \u2014 plus final delivery in Week 12. "
    "Bars show tasks grouped by phase; the red diamonds are milestones: client "
    "presentations and formal sign-offs in Weeks 7, 9 and 11, and the project handover "
    "with all assessment submissions in Week 12.")

add_para(doc2,
    "Team: Bipin Pija (leader \u2013 sprint planning, architecture, client communication, "
    "deployment oversight), Jiban Thapa Magar (deputy \u2013 documentation, requirement "
    "finalization, test logs), Sushil Gaihre (UI/UX \u2013 wireframes, mockups, 360\u00b0 scene "
    "design, frontend UI) and Safal Sharma (developer \u2013 frontend build, backend API, "
    "database schema, core system implementation).")

p = doc2.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture("gantt-team.png", width=Inches(6.9))

doc2.save("Team-Project-Plan-Gantt.docx")
print("wrote Team-Project-Plan-Gantt.docx")
